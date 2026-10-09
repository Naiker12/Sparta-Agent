import json
import threading
from pathlib import Path

import pytest
from core.documents.generation import generate_document
from core.documents.artifact_validation import validate_generated_artifact


def created_path(result, directory):
    assert not result.startswith("Error:"), result
    files = json.loads(result.split("\n__FILES__:")[1])
    path = directory / files[0]["name"]
    assert files[0]["size"] == path.stat().st_size
    assert validate_generated_artifact(str(path)).valid
    return path


def test_pdf_has_real_content_and_paginates(tmp_path):
    import pymupdf
    text = "Informe de España: ventas y análisis.\n\n" * 180
    path = created_path(generate_document({"filename": "informe.pdf", "title": "Resumen", "content": text}, tmp_path), tmp_path)
    with pymupdf.open(path) as document:
        assert len(document) > 1
        assert "Resumen" in document[0].get_text()
        assert "España" in "".join(page.get_text() for page in document)


def test_excel_keeps_numbers_and_literal_cells(tmp_path):
    from openpyxl import load_workbook
    path = created_path(generate_document({"filename": "ventas.xlsx", "headers": ["Producto", "Total"], "rows": [["Café", 125.5], ["=not-a-formula", 42]]}, tmp_path), tmp_path)
    workbook = load_workbook(path)
    assert workbook.active["B2"].value == 125.5
    assert workbook.active["A3"].data_type == "s"
    assert workbook.active.freeze_panes == "A2"
    workbook.close()


def test_word_contains_text_and_table(tmp_path):
    from docx import Document
    path = created_path(generate_document({"filename": "resumen.docx", "title": "Resumen", "content": "Texto final", "headers": ["Nombre"], "rows": [["Ana"]]}, tmp_path), tmp_path)
    document = Document(path)
    assert "Texto final" in [paragraph.text for paragraph in document.paragraphs]
    assert document.tables[0].cell(1, 0).text == "Ana"


@pytest.mark.parametrize("filename", ["../outside.pdf", "C:\\outside.xlsx", "evil.py", "", ".hidden.pdf"])
def test_rejects_paths_and_unsupported_formats(tmp_path, filename):
    assert generate_document({"filename": filename, "content": "Test"}, tmp_path).startswith("Error:")
    assert list(tmp_path.iterdir()) == []


def test_existing_file_is_preserved_and_scratch_is_removed(tmp_path):
    (tmp_path / "notes.txt").write_text("Existing")
    path = created_path(generate_document({"filename": "notes.txt", "content": "New"}, tmp_path), tmp_path)
    assert path.name == "notes (1).txt"
    assert (tmp_path / "notes.txt").read_text() == "Existing"
    assert not list(tmp_path.glob(".sparta-document-*"))


def test_cancelled_generation_has_no_download(tmp_path):
    cancellation = threading.Event()
    cancellation.set()
    assert generate_document({"filename": "cancelled.pdf", "content": "Test"}, tmp_path, cancellation).startswith("Error:")
    assert list(tmp_path.iterdir()) == []


def test_execution_does_not_use_python_tool(tmp_path, monkeypatch):
    from core.inference import tools
    from core.inference.tool_loop_controller import strip_result_for_model
    monkeypatch.setattr(tools, "get_sandbox_workdir", lambda session_id: str(tmp_path))
    monkeypatch.setattr(tools, "_python_exec", lambda *args, **kwargs: pytest.fail("Document tool must not execute Python code"))
    result = tools.execute_tool("generate_document", {"filename": "datos.csv", "headers": ["Valor"], "rows": [[7]]}, session_id="document-test")
    path = created_path(result, tmp_path)
    assert "7" in path.read_text(encoding="utf-8-sig")
    assert "__FILES__" not in strip_result_for_model(result, "generate_document")
