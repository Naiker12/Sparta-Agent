"""Exercise readers without the source tree or any RAG engine present."""
import shutil
import subprocess
import sys
from pathlib import Path


def test_document_readers_from_minimal_resources(tmp_path):
    import pymupdf
    import docx
    import openpyxl

    source = Path(__file__).resolve().parents[1]
    packaged = tmp_path / 'resources' / 'backend'
    readers = packaged / 'core' / 'rag'
    readers.mkdir(parents=True)
    for name in ('__init__.py', 'config.py', 'parsers.py'):
        shutil.copy2(source / 'core' / 'rag' / name, readers / name)
    files = tmp_path / 'documents'
    files.mkdir()
    (files / 'sample.txt').write_text('Spartan TXT marker', encoding='utf-8')
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), 'Spartan PDF marker')
        pdf.save(files / 'sample.pdf')
    word = docx.Document()
    word.add_paragraph('Spartan DOCX marker')
    word.save(files / 'sample.docx')
    book = openpyxl.Workbook()
    book.active.append(['Spartan XLSX marker'])
    book.save(files / 'sample.xlsx')
    book.close()
    program = '''
import sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from core.rag.parsers import parse
assert not any(x in sys.modules for x in ('torch', 'transformers', 'openpyxl', 'docx', 'pymupdf'))
for extension in ('txt', 'pdf', 'docx', 'xlsx'):
    text = '\\n'.join(p.text for p in parse(str(Path(sys.argv[2]) / ('sample.' + extension))))
    assert ('Spartan ' + extension.upper() + ' marker') in text, (extension, text)
assert 'torch' not in sys.modules
assert 'transformers' not in sys.modules
'''
    result = subprocess.run(
        [sys.executable, '-I', '-c', program, str(packaged), str(files)],
        cwd=tmp_path, capture_output=True, text=True, timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
