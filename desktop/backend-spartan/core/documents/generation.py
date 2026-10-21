"""Create document artifacts from data without executing model-supplied code."""

from __future__ import annotations

import csv
import html
import json
import os
import shutil
import tempfile
from pathlib import Path

from core.documents.artifact_validation import validate_generated_artifact

SOURCE_EXTENSIONS = frozenset({".html", ".htm", ".css", ".js", ".mjs", ".ts", ".tsx", ".jsx", ".json", ".yaml", ".yml", ".xml", ".svg", ".py", ".sql", ".sh", ".ps1"})
_RESERVED_BASENAMES = frozenset({"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))})

GENERATE_DOCUMENT_TOOL = {
    "type": "function",
    "function": {
        "name": "generate_document",
        "description": "Create a downloadable document (PDF, XLSX, DOCX, CSV, TXT, Markdown) or source file (HTML, CSS, JS, MJS, TS, TSX, JSX, JSON, YAML, XML, SVG, PY, SQL, SH, PS1) in the active file storage without executing code. For source files pass the exact complete source in content, without Markdown fences, title or table data. For a single HTML page embed CSS and JavaScript. This works independently of the Code toggle: do not use Python or terminal just to save source text. A folder path in the message does not select the output folder. Prefer write_file for source files in an authorized connected workspace; create_file only creates empty files/directories. Do not claim a file exists until the tool succeeds. The application renders the download card; do not invent sandbox links or local download URLs.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Filename with a supported document or source extension, e.g. calculator.html or styles.css. No directory paths."},
                "title": {"type": "string", "description": "Document title."},
                "content": {"type": "string", "description": "Finished document text, or exact source code for a source file. Source code is saved as UTF-8, never executed."},
                "headers": {"type": "array", "items": {"type": "string"}, "description": "Optional table column names."},
                "rows": {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number", "boolean", "null"]}}, "description": "Table or spreadsheet data, one array per row."},
            },
            "required": ["filename"],
        },
    },
}


def _validated_data(arguments):
    filename = arguments.get("filename", "")
    if not isinstance(filename, str) or not filename.strip() or len(filename) > 160:
        raise ValueError("A filename of at most 160 characters is required.")
    filename = filename.strip()
    if filename.startswith(".") or filename.endswith((".", " ")) or filename.split(".")[0].upper() in _RESERVED_BASENAMES or any(char in filename for char in '/\\:<>"|?*') or any(ord(char) < 32 for char in filename):
        raise ValueError("Use a filename without directory paths or reserved characters.")
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".xlsx", ".docx", ".csv", ".txt", ".md", *SOURCE_EXTENSIONS}:
        raise ValueError("Unsupported file extension. Use a supported document or source file format.")
    title = arguments.get("title", "")
    content = arguments.get("content", "")
    headers = arguments.get("headers", [])
    rows = arguments.get("rows", [])
    if not isinstance(title, str) or not isinstance(content, str) or len(title) > 500 or len(content) > 100_000:
        raise ValueError("Document text exceeds the supported limit.")
    if not isinstance(headers, list) or len(headers) > 100 or any(not isinstance(value, str) for value in headers):
        raise ValueError("Headers must contain at most 100 column names.")
    if not isinstance(rows, list) or len(rows) > 10_000:
        raise ValueError("Tables support at most 10,000 rows.")
    for row in rows:
        if not isinstance(row, list) or len(row) > 100 or any(value is not None and not isinstance(value, (str, int, float, bool)) for value in row):
            raise ValueError("Each row must contain at most 100 text, number or boolean cells.")
    if len(json.dumps([headers, rows], ensure_ascii=False)) > 2_000_000:
        raise ValueError("Table data exceeds the supported limit.")
    populated_table = any(str(value).strip() for row in [headers, *rows] for value in row if value is not None)
    if not (content.strip() or populated_table):
        raise ValueError("Prepare the finished document content or populated table data before creating the file; a title alone is not enough.")
    if extension in {".xlsx", ".csv"} and not populated_table:
        raise ValueError("Provide headers or rows for a spreadsheet.")
    if extension in SOURCE_EXTENSIONS and (not content.strip() or headers or rows):
        raise ValueError("Source files require complete source text in content and no table data.")
    return filename, title, content, headers, rows


def _render(path, title, content, headers, rows, cancel_event):
    extension = path.suffix.lower()
    table = ([headers] if headers else []) + rows
    if extension in SOURCE_EXTENSIONS:
        path.write_text(content, encoding="utf-8", newline="")
    elif extension == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Datos"
        for row in table:
            sheet.append(row)
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if headers:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="263645")
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        workbook.save(path)
        workbook.close()
    elif extension == ".docx":
        from docx import Document
        document = Document()
        if title:
            document.add_heading(title, 0)
        for paragraph in content.split("\n\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)
        if table:
            width = max(len(row) for row in table)
            if width:
                grid = document.add_table(rows=len(table), cols=width)
                grid.style = "Table Grid"
                for cells, values in zip(grid.rows, table):
                    for cell, value in zip(cells.cells, values):
                        cell.text = "" if value is None else str(value)
        document.save(path)
    elif extension == ".pdf":
        import pymupdf
        body = f"<h1>{html.escape(title)}</h1>" if title else ""
        body += "".join(f"<p>{html.escape(paragraph).replace(chr(10), '<br/>')}</p>" for paragraph in content.split("\n\n") if paragraph.strip())
        if table:
            body += "<table>" + "".join("<tr>" + "".join(f"<td>{html.escape('' if value is None else str(value))}</td>" for value in row) + "</tr>" for row in table) + "</table>"
        story = pymupdf.Story(body, user_css="body { font-family: sans-serif; font-size: 11pt; } h1 { font-size: 22pt; color: #263645; } p { margin-bottom: 10pt; } td { padding: 5pt; border-bottom: 1px solid #cccccc; } table { width: 100%; }")
        writer = pymupdf.DocumentWriter(str(path))
        def page_rect(number, filled):
            if number >= 200:
                raise ValueError("PDF exceeds the 200 page limit.")
            if cancel_event is not None and cancel_event.is_set():
                raise ValueError("Document generation cancelled.")
            page = pymupdf.paper_rect("a4")
            return page, page + (40, 40, -40, -40), None
        try:
            story.write(writer, page_rect)
        finally:
            writer.close()
    elif extension == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(table)
    else:
        text = (title + "\n\n" if title else "") + content
        if table:
            text += "\n\n" + "\n".join("\t".join("" if value is None else str(value) for value in row) for row in table)
        path.write_text(text, encoding="utf-8")


def generate_document(arguments, workdir, cancel_event=None):
    temporary = None
    try:
        filename, title, content, headers, rows = _validated_data(arguments)
        if cancel_event is not None and cancel_event.is_set():
            raise ValueError("Document generation cancelled.")
        directory = Path(workdir)
        descriptor, name = tempfile.mkstemp(prefix=".sparta-document-", suffix=Path(filename).suffix.lower(), dir=directory)
        os.close(descriptor)
        temporary = Path(name)
        _render(temporary, title, content, headers, rows, cancel_event)
        validation = validate_generated_artifact(str(temporary))
        if not validation.valid:
            raise ValueError(validation.reason or "Generated document is invalid.")
        if cancel_event is not None and cancel_event.is_set():
            raise ValueError("Document generation cancelled.")
        for number in range(1000):
            candidate = directory / (filename if number == 0 else f"{Path(filename).stem} ({number}){Path(filename).suffix}")
            try:
                output = candidate.open("xb")
            except FileExistsError:
                continue
            try:
                with output, temporary.open("rb") as source:
                    shutil.copyfileobj(source, output)
            except Exception:
                candidate.unlink(missing_ok=True)
                raise
            entries = [{"name": candidate.name, "size": candidate.stat().st_size}]
            return f"Created {candidate.name} in {directory}.\n__FILES__:{json.dumps(entries, ensure_ascii=False)}"
        raise ValueError("Too many documents with the same filename.")
    except ImportError as error:
        return f"Error: Document generation dependency unavailable: {error.name}. Update the desktop backend runtime."
    except (ValueError, OSError, TypeError) as error:
        return f"Error: {error}"
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
