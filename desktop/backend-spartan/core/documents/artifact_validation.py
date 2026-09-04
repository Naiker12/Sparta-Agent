"""Bounded validation for document artifacts before the UI advertises them.

The sandbox may create any filename.  A ``.xlsx`` suffix is therefore not proof
that the bytes form a workbook.  This module deliberately validates only the
formats for which the application can make a reliable local assertion; all
other files remain ordinary sandbox downloads.
"""

from __future__ import annotations

import os
import zipfile
from dataclasses import dataclass


_MAX_DOCUMENT_BYTES = 200 * 1024 * 1024
_MAX_ZIP_ENTRIES = 10_000
_MAX_ZIP_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
_MAX_ZIP_EXPANSION_RATIO = 100

_OFFICE_REQUIRED_MEMBERS: dict[str, frozenset[str]] = {
    ".xlsx": frozenset({"[Content_Types].xml", "xl/workbook.xml"}),
    ".xlsm": frozenset({"[Content_Types].xml", "xl/workbook.xml"}),
    ".xltx": frozenset({"[Content_Types].xml", "xl/workbook.xml"}),
    ".xltm": frozenset({"[Content_Types].xml", "xl/workbook.xml"}),
    ".docx": frozenset({"[Content_Types].xml", "word/document.xml"}),
}


@dataclass(frozen = True)
class ArtifactValidation:
    valid: bool
    reason: str | None = None


def _invalid(reason: str) -> ArtifactValidation:
    return ArtifactValidation(False, reason)


def _check_office_package(path: str, required_members: frozenset[str]) -> ArtifactValidation:
    try:
        with zipfile.ZipFile(path) as package:
            members = package.infolist()
            if len(members) > _MAX_ZIP_ENTRIES:
                return _invalid("office package has too many entries")
            names = {member.filename for member in members}
            if not required_members.issubset(names):
                return _invalid("office package is missing required document parts")

            uncompressed = sum(member.file_size for member in members)
            compressed = sum(member.compress_size for member in members)
            if uncompressed > _MAX_ZIP_UNCOMPRESSED_BYTES:
                return _invalid("office package expands beyond the safe limit")
            if compressed and uncompressed > compressed * _MAX_ZIP_EXPANSION_RATIO:
                return _invalid("office package compression ratio is unsafe")
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile):
        return _invalid("file is not a valid Office document package")
    return ArtifactValidation(True)


def _validate_spreadsheet(path: str) -> ArtifactValidation:
    package = _check_office_package(path, _OFFICE_REQUIRED_MEMBERS[os.path.splitext(path)[1].lower()])
    if not package.valid:
        return package
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only = True, data_only = False)
        try:
            if not workbook.sheetnames:
                return _invalid("workbook has no worksheets")
        finally:
            workbook.close()
    except Exception:
        return _invalid("workbook could not be opened")
    return ArtifactValidation(True)


def _validate_word_document(path: str) -> ArtifactValidation:
    package = _check_office_package(path, _OFFICE_REQUIRED_MEMBERS[".docx"])
    if not package.valid:
        return package
    try:
        from docx import Document

        Document(path)
    except Exception:
        return _invalid("Word document could not be opened")
    return ArtifactValidation(True)


def _validate_pdf(path: str) -> ArtifactValidation:
    try:
        with open(path, "rb") as handle:
            if handle.read(5) != b"%PDF-":
                return _invalid("file does not start with a PDF header")
        import fitz

        document = fitz.open(path)
        try:
            if document.page_count < 1:
                return _invalid("PDF has no pages")
        finally:
            document.close()
    except Exception:
        return _invalid("PDF could not be opened")
    return ArtifactValidation(True)


def validate_generated_artifact(path: str) -> ArtifactValidation:
    """Validate a supported document artifact without trusting its extension.

    Unknown file types intentionally return valid: this is an artifact guard,
    not a general file-type policy.  Document validation is bounded before any
    parser receives the file, so a malformed archive cannot turn a file card
    into an unbounded parsing operation.
    """
    extension = os.path.splitext(path)[1].lower()
    if extension not in {*_OFFICE_REQUIRED_MEMBERS, ".pdf"}:
        return ArtifactValidation(True)
    try:
        size = os.path.getsize(path)
    except OSError:
        return _invalid("file disappeared before it could be validated")
    if size <= 0:
        return _invalid("document is empty")
    if size > _MAX_DOCUMENT_BYTES:
        return _invalid("document exceeds the 200 MB validation limit")
    if extension == ".pdf":
        return _validate_pdf(path)
    if extension == ".docx":
        return _validate_word_document(path)
    return _validate_spreadsheet(path)
